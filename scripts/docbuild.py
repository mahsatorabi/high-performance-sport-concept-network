"""Shared helpers for building the Article 1 manuscript."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"D:\Uni of Birjand\articles\adel\article1\Article1_manuscript.docx"


def new_doc():
    d = Document()
    s = d.styles["Normal"]
    s.font.name = "Times New Roman"
    s.font.size = Pt(12)
    s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = s.paragraph_format
    pf.line_spacing = 2.0
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for sec in d.sections:
        sec.top_margin = sec.bottom_margin = Inches(1.0)
        sec.left_margin = sec.right_margin = Inches(1.0)
    return d


def _style_font(style, name, size, bold=False, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        style.font.color.rgb = color


def h1(d, t):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(12)
    r.font.name = "Times New Roman"
    return p


def h2(d, t):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(t)
    r.bold = True
    r.italic = True
    r.font.size = Pt(12)
    return p


def title(d, t):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(16)
    return p


def para(d, t, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, italic=False,
         space_after=0):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    r = p.add_run(t)
    r.font.size = Pt(size)
    r.italic = italic
    return p


def rich(d, parts, indent=True, size=12, space_after=0):
    """parts = list of (text, italic) tuples."""
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.27)
    for txt, it in parts:
        r = p.add_run(txt)
        r.font.size = Pt(size)
        r.italic = it
    return p


def bullet(d, t, size=11):
    p = d.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t)
    r.font.size = Pt(size)
    return p


def table_caption(d, t):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(10)
    return p


def figure_caption(d, t):
    p = d.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(t)
    r.bold = True
    r.font.size = Pt(10)
    return p


def table(d, header, rows, widths=None, font=8.5):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htxt in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(htxt))
        r.bold = True
        r.font.size = Pt(font)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E8EDF2")
        hdr[i]._tc.get_or_add_tcPr().append(shade)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(v))
            r.font.size = Pt(font)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def figure(d, path, width_in=6.0):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Inches(width_in))
    return p


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def rule(d):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "999999")
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p